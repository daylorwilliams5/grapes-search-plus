"""
make_agent_updates.py
=====================
Builds a spreadsheet of what the research agent found, for entering into the
GRAPES admin system. The agent's changes went into the search site's data;
the admin system still has the values from before the first batch.

"Before" is the site data at --since (the commit before the first agent
batch was merged), which came from the DGE spreadsheet. "After" is the
current data. Each change carries the quote and source page the agent found,
taken from the batch reports on the agent-research/* branches.

Usage:
    python scripts/make_agent_updates.py [--since 6f255d8] [--out worklists/...xlsx]

Requirements:
    pip install openpyxl
"""
import argparse, collections, datetime, json, re, subprocess

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ap = argparse.ArgumentParser()
ap.add_argument("--since", default="6f255d8", help="Commit with the data before the agent's first batch")
ap.add_argument("--out")
args = ap.parse_args()
TODAY = datetime.date.today()
out_path = args.out or f"worklists/grapes_agent_updates_{TODAY:%Y-%m}.xlsx"


def git(*a):
    return subprocess.check_output(["git", *a]).decode()


before = {r["id"]: r for r in json.loads(git("show", f"{args.since}:data/fellowships_master.json"))}
now = json.load(open("data/fellowships_master.json"))

# Evidence from every merged batch report, latest batch last so it wins
merged_log = git("log", "--format=%s%n%b", "origin/main")
evidence = {}
branches = sorted(b.strip().removeprefix("origin/") for b in git("branch", "-r").splitlines()
                  if "agent-research/" in b)
for br in branches:
    if br not in merged_log:
        continue
    try:
        report = json.loads(git("show", f"origin/{br}:data/agent_proposals.json"))
    except subprocess.CalledProcessError:
        continue
    for res in report["results"]:
        if not res.get("error"):
            evidence[res["id"]] = res


def day(s):
    try:
        d = datetime.date.fromisoformat(s or "")
        return None if d.year == 9999 else d
    except ValueError:
        return None


def site_url(u):
    return re.sub(r"^https?://", "", (u or "").strip()).rstrip("/")


GENERIC = {"fellowship", "fellowships", "program", "programs", "graduate", "research", "postdoctoral",
           "doctoral", "predoctoral", "award", "awards", "grant", "grants", "scholarship", "scholarships",
           "dissertation", "student", "students", "foundation", "national", "american", "international",
           "center", "institute", "society", "association", "for", "the", "and", "with", "from", "study",
           "studies", "summer", "travel", "early", "career", "science", "sciences", "university", "ucla"}
def words(t):
    return re.findall(r"[A-Za-z][A-Za-z&.-]*[A-Za-z]", t or "")
freq = collections.Counter(w.lower() for r in now for w in set(words(r["title"])))
def search_word(r):
    c = [w for w in words(r["title"]) if w.lower() not in GENERIC and len(w) >= 3] or \
        [w for w in words(r.get("agency1")) if w.lower() not in GENERIC and len(w) >= 3]
    return min(c, key=lambda w: (freq.get(w.lower(), 0), -len(w))) if c else ""


def field_evidence(rid, name, value_matches):
    """(source, quote, year added) for a field, if the agent's finding is what's now on file."""
    c = evidence.get(rid, {}).get("checks", {}).get(name, {})
    if c.get("verdict") != "verified" or not value_matches(c.get("value", "")):
        return "", "", ""
    return c.get("source_url", ""), c.get("quote", ""), "Yes" if c.get("year_inferred") else ""


# ---------- Sort the findings ----------
updates, confirmed, attention = [], [], []
for r in now:
    o = before.get(r["id"])
    if not o:
        continue
    new_dl = day(r["deadlineSort"]) if r.get("deadline") != o.get("deadline") else None
    new_amt = r.get("amount") if r.get("amount") != o.get("amount") else ""
    new_url = site_url(r.get("officialUrl")) if r.get("officialUrl") != o.get("officialUrl") else ""
    if new_dl or new_amt or new_url:
        src, quote, yr = field_evidence(r["id"], "deadline",
                                        lambda v: new_dl and v == new_dl.isoformat())
        if not src and new_amt:
            src, quote, _ = field_evidence(r["id"], "amount", lambda v: v == new_amt)
        updates.append(dict(r=r, o=o, dl=new_dl, amt=new_amt, url=new_url, src=src, quote=quote, yr=yr))
    elif r.get("checked") and r.get("status") != "Published":
        src, quote, yr = field_evidence(r["id"], "deadline", lambda v: v == r["deadlineSort"])
        confirmed.append(dict(r=r, src=src or r.get("checkedSource", ""), quote=quote, yr=yr))

# Listings staff need to look at: programs the agent found ended or paused,
# and citizenship rules that conflict with the listing
for rid, res in evidence.items():
    r = next((x for x in now if x["id"] == rid), None)
    reasons = [f for f in res.get("flags", []) if "discontinued" in f or "paused" in f or "Citizenship" in f]
    if r and reasons:
        attention.append(dict(r=r, reasons="; ".join(reasons), notes=res.get("notes") or ""))


def norm_url(u):
    return re.sub(r"^https?://(www\.)?", "", (u or "").strip()).rstrip("/").lower()


# Still unverified after the agent's check, with its explanation (this is
# where broken links and unposted deadlines show up)
TODAY_ISO = TODAY.isoformat()
def verified(r):
    return r.get("status") == "Published" or (
        r.get("checked") and (TODAY - datetime.date.fromisoformat(r["checked"])).days <= 183)
unverified = [dict(r=r, notes=(evidence.get(r["id"]) or {}).get("notes", "")) for r in now if not verified(r)]
unverified.sort(key=lambda u: u["r"]["id"])

# The agent confirmed the program on a different page than the one on file
other_pages = []
for rid, res in evidence.items():
    r = next((x for x in now if x["id"] == rid), None)
    ou = res.get("checks", {}).get("official_url", {})
    if r and ou.get("verdict") == "verified" and norm_url(ou.get("value")) != norm_url(r.get("officialUrl")):
        other_pages.append(dict(r=r, url=ou["value"], quote=ou.get("quote", "")))
other_pages.sort(key=lambda o: o["r"]["id"])

# Dates the agent got wrong or the two runs disagreed on, found in review
MANUAL = {72: "Sponsor page gives Dec 1, 2027 (likely a typo); confirm the date",
          135: "Sponsor page has moved to the following cycle; confirm whether Nov 1, 2026 still applies",
          360: "Two agent runs disagreed on the year; check the deadline",
          519: "Two agent runs disagreed on the year; check the deadline",
          66: "Sponsor states $30,000, not $27,500; update the amount",
          508: "Agent filled the amount with a note, not an amount; corrected on the site to the salary grade",
          1049: "DGE spreadsheet uses 1049 for both Clark Predoctoral and Ahmanson-Getty; fix the spreadsheet"}
have = {a["r"]["id"] for a in attention}
for rid, why in MANUAL.items():
    r = next((x for x in now if x["id"] == rid), None)
    if r and rid not in have:
        attention.append(dict(r=r, reasons=why, notes=(evidence.get(rid) or {}).get("notes", "")))
    elif r:
        a = next(a for a in attention if a["r"]["id"] == rid)
        a["reasons"] = why + "; " + a["reasons"]

updates.sort(key=lambda u: (u["dl"] is None or u["dl"] < TODAY, u["dl"] or datetime.date.max))
confirmed.sort(key=lambda c: c["r"]["deadlineSort"])
attention.sort(key=lambda a: (not any(w in a["reasons"] for w in ("discontinued", "paused")), a["r"]["id"]))

# ---------- Workbook ----------
F = "Arial"
base, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF")
hdr_fill, in_hdr_fill = PatternFill("solid", fgColor="2D68C4"), PatternFill("solid", fgColor="B7791F")
in_fill, new_fill = PatternFill("solid", fgColor="FFF7D6"), PatternFill("solid", fgColor="E3EEFB")
border = Border(bottom=Side(style="thin", color="D0D0D0"))
link_font = Font(name=F, size=10, color="0563C1", underline="single")


def sheet(wb, title, cols, rows, first_input, new_cols=(), wrap=(), dates=(), links=(), done=None):
    ws = wb.create_sheet(title)
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.alignment = hdr_font, Alignment(wrap_text=True, vertical="center")
        c.fill = in_hdr_fill if j >= first_input else hdr_fill
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 32
    for i, vals in enumerate(rows, 2):
        for j, v in enumerate(vals, 1):
            h = cols[j - 1][0]
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = base, border
            c.alignment = Alignment(vertical="top", wrap_text=h in wrap)
            if h in dates:
                c.number_format = "m/d/yyyy"
            if h in new_cols and v not in (None, ""):
                c.fill, c.font = new_fill, bold
            if j >= first_input:
                c.fill = in_fill
            if h in links and v:
                c.hyperlink = v if str(v).startswith("http") else "https://" + str(v)
                c.font = link_font
    n = len(rows) + 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(n, 2)}"
    if done:
        col = get_column_letter(first_input + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(done) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(n, 2)}")
        ws.conditional_formatting.add(f"A2:{col}{max(n, 2)}", FormulaRule(
            formula=[f'${col}2<>""'], fill=PatternFill("solid", fgColor="E8F3EA")))
    return ws


wb = Workbook()
ws0 = wb.active
ws0.title = "How to use"
ws0.sheet_view.showGridLines = False
lines = [
    ("GRAPES updates from the research agent", Font(name=F, size=14, bold=True)),
    (f"Built {TODAY:%B %d, %Y}. The agent checked all {len(now)} listings against sponsors' websites, Oct 15-19, 2026.", base),
    ("Its changes are on the public search site, but the GRAPES admin system still has the old values.", base),
    ("", base),
    ("Tabs", bold),
    (f"Update in GRAPES: {len(updates)} records with a new deadline, amount, or website. Upcoming deadlines first.", base),
    (f"Mark verified: {len(confirmed)} unverified records whose current deadline the agent confirmed. Nothing to", base),
    ("   change except checking the deadline matches and marking the record published.", base),
    (f"Needs a person: {len(attention)} records to decide on (ended or paused programs, odd dates, conflicts).", base),
    (f"Still unverified: {len(unverified)} records the agent couldn't confirm, with its explanation (often the sponsor", base),
    ("   hasn't posted the next deadline yet, or the link on file is broken or points to the wrong page).", base),
    (f"Different page found: {len(other_pages)} records where the agent confirmed the program on a different page than", base),
    ("   the one on file. Switch the Website field if the new page is better.", base),
    ("", base),
    ("For each row", bold),
    ("1. Search the admin form for the Search word and open the record with the matching Record Number.", base),
    ("2. Enter the blue values. Deadlines are M/D/YYYY and websites have no https://, as the form asks.", base),
    ("3. If you want to double-check, open the Source page: the Quote column is copied from it.", base),
    ("4. Pick a status in Done. The row turns green.", base),
    ("", base),
    ("\"Year added\" means the sponsor's page gave the date without a year (e.g. \"due March 1 each year\")", base),
    ("and the agent filled in the next one.", base),
]
for i, (t, f) in enumerate(lines, 1):
    ws0.cell(row=i, column=1, value=t).font = f
ws0.column_dimensions["A"].width = 120

DONE = ["Entered", "Already up to date", "Sponsor page differs", "Record not found"]
cols = [("Record #", 9), ("Search word", 14), ("Title", 40), ("GRAPES had: deadline", 13),
        ("New deadline", 13), ("Year added", 8), ("New amount", 30), ("New website", 34),
        ("Source page", 30), ("Quote from the page", 44), ("Backend shows", 14), ("Done", 18), ("Notes", 30)]
rows = [[u["r"]["id"], search_word(u["r"]), u["r"]["title"], day(u["o"]["deadlineSort"]),
         u["dl"], u["yr"], u["amt"], u["url"], u["src"], u["quote"], None, None, None] for u in updates]
sheet(wb, "Update in GRAPES", cols, rows, first_input=11,
      new_cols={"New deadline", "New amount", "New website"},
      wrap={"Title", "New amount", "New website", "Source page", "Quote from the page", "Notes"},
      dates={"GRAPES had: deadline", "New deadline"}, links={"Source page", "New website"}, done=DONE)

cols = [("Record #", 9), ("Search word", 14), ("Title", 40), ("Confirmed deadline", 13), ("Year added", 8),
        ("Source page", 34), ("Quote from the page", 50), ("Backend shows", 14), ("Done", 22), ("Notes", 30)]
rows = [[c["r"]["id"], search_word(c["r"]), c["r"]["title"], day(c["r"]["deadlineSort"]), c["yr"],
         c["src"], c["quote"], None, None, None] for c in confirmed]
sheet(wb, "Mark verified", cols, rows, first_input=8, new_cols={"Confirmed deadline"},
      wrap={"Title", "Source page", "Quote from the page", "Notes"}, dates={"Confirmed deadline"},
      links={"Source page"}, done=["Marked published", "Deadline differs", "Record not found"])

cols = [("Record #", 9), ("Search word", 14), ("Title", 40), ("Why", 40), ("Agent's notes", 70),
        ("Decision", 22), ("Notes", 30)]
rows = [[a["r"]["id"], search_word(a["r"]), a["r"]["title"], a["reasons"], (a["notes"] or "")[:600],
         None, None] for a in attention]
sheet(wb, "Needs a person", cols, rows, first_input=6, wrap={"Title", "Why", "Agent's notes", "Notes"},
      done=None)
ws = wb["Needs a person"]
dv = DataValidation(type="list", formula1='"Fixed,Remove from GRAPES,Unpublish for now,Ask sponsor,No change"',
                    allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"F2:F{len(attention) + 1}")

cols = [("Record #", 9), ("Search word", 14), ("Title", 40), ("Deadline on file", 13),
        ("Website on file", 34), ("Why the agent couldn't confirm it", 80), ("Decision", 22), ("Notes", 30)]
rows = [[u["r"]["id"], search_word(u["r"]), u["r"]["title"], day(u["r"]["deadlineSort"]),
         re.sub(r"^https?://", "", u["r"].get("officialUrl") or ""), (u["notes"] or "")[:700], None, None]
        for u in unverified]
sheet(wb, "Still unverified", cols, rows, first_input=7, wrap={"Title", "Website on file", "Why the agent couldn't confirm it", "Notes"},
      dates={"Deadline on file"}, links={"Website on file"},
      done=None)
ws = wb["Still unverified"]
dv = DataValidation(type="list", formula1='"Checked by hand,Fixed link,Waiting for sponsor,Remove from GRAPES"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"G2:G{len(unverified) + 1}")

cols = [("Record #", 9), ("Search word", 14), ("Title", 40), ("Website on file", 40),
        ("Page the agent confirmed", 40), ("Quote from that page", 44), ("Decision", 18), ("Notes", 30)]
rows = [[o["r"]["id"], search_word(o["r"]), o["r"]["title"], re.sub(r"^https?://", "", o["r"].get("officialUrl") or ""),
         re.sub(r"^https?://", "", o["url"]), o["quote"], None, None] for o in other_pages]
sheet(wb, "Different page found", cols, rows, first_input=7,
      new_cols={"Page the agent confirmed"}, wrap={"Title", "Website on file", "Page the agent confirmed", "Quote from that page", "Notes"},
      links={"Website on file", "Page the agent confirmed"})
ws = wb["Different page found"]
dv = DataValidation(type="list", formula1='"Switched,Keep current"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"G2:G{len(other_pages) + 1}")

wb.save(out_path)
print(f"Saved {out_path}: {len(updates)} updates, {len(confirmed)} to mark verified, "
      f"{len(attention)} need a person, {len(unverified)} still unverified, "
      f"{len(other_pages)} with a different page; evidence from {len(evidence)} listings")
