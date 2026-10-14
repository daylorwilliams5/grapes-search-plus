"""
make_update_list.py
===================
Builds a spreadsheet of specific updates to enter in the GRAPES admin
system: the deadlines and links that were confirmed on sponsors' sites in the
September 2026 audit but differ from what the DGE spreadsheet (and so,
probably, the backend) had.

The "before" values come from the site data just before the audit, which was
built from the DGE spreadsheet. The "after" values are the site's current
data. Every value should still be checked on the sponsor's page before it's
entered, since the backend may have been updated since.

Usage:
    python scripts/make_update_list.py [--out worklists/grapes_updates_to_enter.xlsx]

Requirements:
    pip install openpyxl
"""
import argparse, collections, datetime, json, re, subprocess

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

AUDIT_COMMIT = '0b6f342'  # "Apply September 2026 listing audit"

ap = argparse.ArgumentParser()
ap.add_argument('--out', default='worklists/grapes_updates_to_enter.xlsx')
args = ap.parse_args()
TODAY = datetime.date.today()

before = json.loads(subprocess.check_output(['git', 'show', f'{AUDIT_COMMIT}^:data/fellowships_master.json']))
now = json.load(open('data/fellowships_master.json'))
excluded = json.load(open('data/excluded_ids.json'))['records']
by_id = collections.defaultdict(list)
for r in before:
    by_id[r['id']].append(r)


def day(s):
    try:
        d = datetime.date.fromisoformat(s or '')
        return None if d.year == 9999 else d
    except ValueError:
        return None


def site_url(u):
    return re.sub(r'^https?://', '', (u or '').strip()).rstrip('/')


def same_url(a, b):
    return site_url(a).replace('www.', '') == site_url(b).replace('www.', '')


GENERIC = {'fellowship', 'fellowships', 'program', 'programs', 'graduate', 'research', 'postdoctoral',
           'doctoral', 'predoctoral', 'award', 'awards', 'grant', 'grants', 'scholarship', 'scholarships',
           'dissertation', 'student', 'students', 'foundation', 'national', 'american', 'international',
           'center', 'institute', 'society', 'association', 'for', 'the', 'and', 'with', 'from', 'study',
           'studies', 'summer', 'travel', 'early', 'career', 'science', 'sciences', 'university', 'ucla'}
def words(t):
    return re.findall(r"[A-Za-z][A-Za-z&.-]*[A-Za-z]", t or '')
freq = collections.Counter(w.lower() for r in now for w in set(words(r['title'])))
def search_word(r):
    c = [w for w in words(r['title']) if w.lower() not in GENERIC and len(w) >= 3] or \
        [w for w in words(r.get('agency1')) if w.lower() not in GENERIC and len(w) >= 3]
    return min(c, key=lambda w: (freq.get(w.lower(), 0), -len(w))) if c else ''


def amount(r):
    a = (r.get('amount') or '').strip()
    return a if a and a.lower() not in ('not specified', 'varies') else 'Not listed: check sponsor page'


# ---------- Work out the updates ----------
updates, unmatched = [], []
for r in now:
    olds = by_id.get(r['id'], [])
    if r['id'] == 1049:
        continue  # has its own tab
    if len(olds) != 1:
        # Added after the audit, or the spreadsheet used this number twice:
        # there's no reliable "before" value, so it goes on the check tab
        unmatched.append(r)
        continue
    o = olds[0]
    old_dl, new_dl = day(o['deadlineSort']), day(r['deadlineSort'])
    change = {}
    if old_dl != new_dl:
        if new_dl is None and r.get('season') == 'Open/Rolling':
            change['open'] = True
        elif new_dl is not None:
            change['deadline'] = (old_dl, new_dl)
    if not same_url(o['officialUrl'], r['officialUrl']):
        change['url'] = (site_url(o['officialUrl']), site_url(r['officialUrl']))
    if change:
        updates.append((r, o, change))


def sort_key(item):
    r, _, ch = item
    d = ch.get('deadline', (None, None))[1]
    upcoming = d is not None and d >= TODAY
    return (0 if upcoming else 1, d or datetime.date(9999, 1, 1), r['title'].lower())
updates.sort(key=sort_key)

# ---------- Styles ----------
F = 'Arial'
base, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
hdr_font = Font(name=F, size=10, bold=True, color='FFFFFF')
hdr_fill, in_hdr_fill = PatternFill('solid', fgColor='2D68C4'), PatternFill('solid', fgColor='B7791F')
in_fill = PatternFill('solid', fgColor='FFF7D6')
new_fill = PatternFill('solid', fgColor='E3EEFB')
border = Border(bottom=Side(style='thin', color='D0D0D0'))
link_font = Font(name=F, size=10, color='0563C1', underline='single')
DATE = 'm/d/yyyy'  # the backend's Deadline Date format


def table(ws, cols, rows, first_input, new_cols=(), wrap=(), dates=(), links=()):
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.alignment = hdr_font, Alignment(wrap_text=True, vertical='center')
        c.fill = in_hdr_fill if j >= first_input else hdr_fill
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 32
    for i, vals in enumerate(rows, 2):
        for j, v in enumerate(vals, 1):
            h = cols[j - 1][0]
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = base, border
            c.alignment = Alignment(vertical='top', wrap_text=h in wrap)
            if h in dates:
                c.number_format = DATE
            if h in new_cols and v not in (None, ''):
                c.fill, c.font = new_fill, bold
            if j >= first_input:
                c.fill = in_fill
            if h in links and v:
                c.hyperlink = v if v.startswith('http') else 'https://' + v
                c.font = link_font
    n = len(rows) + 1
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{n}'
    return n


def done_dropdown(ws, col, n, options):
    dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{n}')
    ws.conditional_formatting.add(f'A2:{col}{n}', FormulaRule(
        formula=[f'${col}2<>""'], fill=PatternFill('solid', fgColor='E8F3EA')))


wb = Workbook()

# ---------- How to use ----------
ws0 = wb.active
ws0.title = 'How to use'
ws0.sheet_view.showGridLines = False
counts = collections.Counter()
for _, _, ch in updates:
    for k in ch:
        counts[k] += 1
lines = [
    ('GRAPES updates to enter', Font(name=F, size=14, bold=True)),
    (f'Built {TODAY:%B %d, %Y}.', base),
    ('', base),
    ('What this is', bold),
    ('In September 2026 every listing on GRAPES Search+ was checked against its sponsor\'s website. These are the', base),
    ('places where what was found differs from the DGE GRAPES spreadsheet, so the admin system probably still has the old value.', base),
    ('', base),
    ('Tabs', bold),
    (f'Updates: {len(updates)} records with a new deadline ({counts["deadline"]}), a new website ({counts["url"]}),', base),
    (f'   or that should be marked Open Deadline ({counts["open"]}). Records with upcoming deadlines are first.', base),
    (f'Check: {len(unmatched)} listings added to the site after the audit (or with a shared record number). There is no', base),
    ('   "before" value, so compare the backend with the researched deadline and website.', base),
    (f'Review for removal: {len(excluded)} programs taken off the site as discontinued, paused, or duplicates. Ask before', base),
    ('   changing these in the backend; it is DGE\'s decision.', base),
    ('Record 1049: the fix for the Clark / Ahmanson-Getty mix-up.', base),
    ('', base),
    ('How to use', bold),
    ('1. Search the admin form for the Search word and open the record with the matching Record Number.', base),
    ('2. Fill in "Backend shows" with the deadline the backend has now. If it already matches the new value, pick', base),
    ('   "Already up to date" in Done and move on.', base),
    ('3. Open the sponsor page (the link in the new Website column) and confirm the new value is still right.', base),
    ('4. Enter it in the backend, then pick "Entered" in Done. Use "Sponsor page differs" if the page says something else,', base),
    ('   and put what it says in Notes.', base),
    ('', base),
    ('Blue cells are the new values. Dates are in M/D/YYYY, the format the backend\'s Deadline Date field asks for.', base),
    ('Websites are shown without https://, the way the backend\'s Website field asks for them.', base),
    ('Award amount is what the sponsor page said in September. "Not listed" means it wasn\'t found; check the sponsor page.', base),
    ('Some new deadlines have already passed: the sponsor hadn\'t posted the next cycle in September. Enter the newer date,', base),
    ('but check the page first in case the next cycle is now posted.', base),
]
for i, (t, f) in enumerate(lines, 1):
    ws0.cell(row=i, column=1, value=t).font = f
ws0.column_dimensions['A'].width = 120

# ---------- Updates ----------
ws = wb.create_sheet('Updates')
cols = [('Record #', 9), ('Search word', 14), ('Title', 42), ('What to change', 26),
        ('Spreadsheet deadline', 13), ('New deadline', 13), ('New deadline passed?', 11),
        ('Open Deadline box', 11), ('Spreadsheet website', 34), ('New website', 34),
        ('Award amount', 28), ('Backend shows', 16), ('Done', 20), ('Notes', 34)]
rows = []
for r, o, ch in updates:
    what = []
    if 'deadline' in ch:
        old, new = ch['deadline']
        # A later spreadsheet date that the sponsor never published was a projection
        what.append('Deadline (spreadsheet date may be a projection; check before changing)'
                    if old and new < old and new < TODAY else 'Deadline')
    if 'open' in ch: what.append('Check Open Deadline')
    if 'url' in ch: what.append('Website')
    old_dl, new_dl = ch.get('deadline', (day(o['deadlineSort']), None))
    rows.append([
        r['id'], search_word(r), r['title'], ', '.join(what),
        old_dl, new_dl, ('Yes' if new_dl and new_dl < TODAY else '') if new_dl else '',
        'Check it' if 'open' in ch else '',
        ch['url'][0] if 'url' in ch else '', ch['url'][1] if 'url' in ch else '',
        amount(r), None, None, None,
    ])
n = table(ws, cols, rows, first_input=12,
          new_cols={'New deadline', 'Open Deadline box', 'New website'},
          wrap={'Title', 'Spreadsheet website', 'New website', 'Award amount', 'Notes'},
          dates={'Spreadsheet deadline', 'New deadline'}, links={'New website'})
done_dropdown(ws, 'M', n, ['Entered', 'Already up to date', 'Sponsor page differs', 'Record not found'])
# Add the current website to every row, so the sponsor page is one click away
for i, (r, _, ch) in enumerate(updates, 2):
    if 'url' not in ch:
        c = ws.cell(row=i, column=10, value=site_url(r['officialUrl']))
        c.hyperlink, c.font = r['officialUrl'], Font(name=F, size=9, color='6B7280', underline='single')
ws.cell(row=1, column=10).value = 'New website (grey = unchanged, link only)'

# ---------- Check ----------
wc = wb.create_sheet('Check')
cols = [('Record #', 9), ('Search word', 14), ('Title', 42), ('Researched deadline', 13),
        ('Rolling?', 9), ('Researched website', 40), ('Award amount', 28), ('Backend shows', 16), ('Done', 20), ('Notes', 34)]
unmatched.sort(key=lambda r: (day(r['deadlineSort']) is None or day(r['deadlineSort']) < TODAY,
                              day(r['deadlineSort']) or datetime.date(9999, 1, 1)))
rows = [[r['id'], search_word(r), r['title'], day(r['deadlineSort']),
         'Yes' if r.get('season') == 'Open/Rolling' and not day(r['deadlineSort']) else '',
         site_url(r['officialUrl']), amount(r), None, None, None] for r in unmatched]
n = table(wc, cols, rows, first_input=8, wrap={'Title', 'Researched website', 'Award amount', 'Notes'},
          dates={'Researched deadline'}, links={'Researched website'})
done_dropdown(wc, 'I', n, ['Matches', 'Entered', 'Sponsor page differs', 'Record not found'])

# ---------- Review for removal ----------
wr = wb.create_sheet('Review for removal')
cols = [('Record #', 9), ('Title', 50), ('Why it was taken off the site', 50), ('DGE decision', 22), ('Notes', 34)]
rows = [[e['id'], e['title'], e['reason'], None, None] for e in sorted(excluded, key=lambda e: e['reason'])]
n = table(wr, cols, rows, first_input=4, wrap={'Title', 'Why it was taken off the site', 'Notes'})
done_dropdown(wr, 'D', n, ['Remove from GRAPES', 'Keep', 'Ask sponsor'])

# ---------- Record 1049 ----------
w9 = wb.create_sheet('Record 1049')
w9.sheet_view.showGridLines = False
items = [
    ('Record 1049: Clark / Ahmanson-Getty', Font(name=F, size=14, bold=True)),
    ('', base),
    ('In the admin system, 1049 is the Ahmanson-Getty Postdoctoral Fellowship. The Clark Predoctoral Fellowships have their own', base),
    ('record. The DGE spreadsheet lists both under 1049, which is why the public site showed Clark and dropped Ahmanson-Getty.', base),
    ('', base),
    ('To do', bold),
    ('1. Admin system, record 1049 (Ahmanson-Getty): change "William Andrew Clark Memorial Library" to "William Andrews Clark', base),
    ('   Memorial Library" in Foundation or Agency Name. The deadline (2/1/2027) is already current.', base),
    ('2. Admin system, Clark Predoctoral Fellowships: write down its Record Number and check its deadline against', base),
    ('   www.1718.ucla.edu/research/fellowships/', base),
    ('3. DGE spreadsheet (ask CJ or whoever keeps it): the "Clark Pre-Doctoral Fellowship" row says 1049. Change it to', base),
    ('   Clark\'s real record number from step 2.', base),
    ('4. Send the Clark record number to Daylor so the public site can be fixed (1049 becomes Ahmanson-Getty again).', base),
    ('', base),
    ('Clark record number:', bold),
]
for i, (t, f) in enumerate(items, 1):
    w9.cell(row=i, column=1, value=t).font = f
w9.cell(row=len(items), column=2).fill = in_fill
w9.column_dimensions['A'].width = 110
w9.column_dimensions['B'].width = 16

wb.save(args.out)
print(f'Saved {args.out}: {len(updates)} updates ({dict(counts)}), {len(unmatched)} to check, '
      f'{len(excluded)} to review for removal')
